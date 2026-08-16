from __future__ import annotations

import math
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT.parent / ".doc-work" / "output"
WEEKLY_REFERENCE = Path(r"C:\Users\ROG\Downloads\PetMind MoE Agent · 工作周报·8.3-8.10.docx")
WEEKLY_OUTPUT = WORK / "PetMind MoE Agent · 工作周报·8.10-8.16.docx"
SOP_OUTPUT = WORK / "PetMind_MoE服务定价与运营SOP调研报告_2026-08-16_v1.2.docx"

FONT = "SimSun"
FONT_EAST_ASIA = "宋体"
MONO = "Consolas"
BLACK = "111111"
INK = "243447"
MUTED = "667085"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LINK_BLUE = "3370FF"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_GOLD = "FFF7DD"
PALE_RED = "FDECEC"
PALE_GREEN = "EDF7F0"
WHITE = "FFFFFF"


def set_run_font(run, *, size=None, bold=None, color=None, italic=None, mono=False) -> None:
    name = MONO if mono else FONT
    east = MONO if mono else FONT_EAST_ASIA
    run.font.name = name
    r_fonts = run._element.get_or_add_rPr().rFonts
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), east)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style(style, *, size, color=INK, bold=False, before=0, after=6, line=1.1) -> None:
    style.font.name = FONT
    fonts = style._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_borders(table, *, color="D7DBE2", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def prevent_row_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def set_table_geometry(table, widths_dxa, *, indent=120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
    if tbl_w.getparent() is None:
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd") or OxmlElement("w:tblInd")
    if tbl_ind.getparent() is None:
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
    if layout.getparent() is None:
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
            if tc_w.getparent() is None:
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths, *, header_fill=LIGHT_GRAY, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_borders(table)
    set_repeat_header(table.rows[0])
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(text))
        set_run_font(run, size=font_size, bold=True, color=DARK_BLUE)
    for row_values in rows:
        row = table.add_row()
        for i, value in enumerate(row_values):
            p = row.cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(str(value))
            set_run_font(run, size=font_size, color=INK)
    set_table_geometry(table, widths)
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(0)
    return table


def add_numbering(doc, *, abstract_id, num_id, fmt, text, left=720, hanging=360) -> None:
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("start", "1"), ("numFmt", fmt), ("lvlText", text), ("lvlJc", "left")):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:val"), value)
        level.append(node)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    p_pr.append(ind)
    level.append(p_pr)
    abstract.append(level)
    index = next((i for i, x in enumerate(numbering) if x.tag == qn("w:num")), len(numbering))
    numbering.insert(index, abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)


def apply_numbering(paragraph, num_id) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)


def add_bullet(doc, text, *, num_id=77, size=11, after=5, color=INK) -> None:
    p = doc.add_paragraph()
    apply_numbering(p, num_id)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    set_run_font(p.add_run(text), size=size, color=color)


def add_number(doc, text, *, num_id=78, size=11) -> None:
    p = doc.add_paragraph()
    apply_numbering(p, num_id)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    set_run_font(p.add_run(text), size=size, color=INK)


def add_body(doc, text, *, bold_prefix=None, indent=False, size=11, after=6) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if indent:
        p.paragraph_format.first_line_indent = Pt(21)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), size=size, bold=True, color=BLACK)
        set_run_font(p.add_run(text[len(bold_prefix):]), size=size, color=INK)
    else:
        set_run_font(p.add_run(text), size=size, color=INK)


def add_callout(doc, label, text, *, fill=PALE_GOLD, width=9360) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(8)
    p.paragraph_format.right_indent = Pt(8)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D7DBE2")
        node.set(qn("w:space"), "6")
        borders.append(node)
    p_pr.append(borders)
    set_run_font(p.add_run(label + "  "), size=10.5, bold=True, color=DARK_BLUE)
    set_run_font(p.add_run(text), size=10.5, color=INK)


def add_field(paragraph, instruction: str, *, size=9, color=MUTED) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=size, color=color)


def set_alt_text(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def font(size, bold=False):
    candidates = [Path(r"C:\Windows\Fonts\simsun.ttc"), Path(r"C:\Windows\Fonts\simsun.ttf")]
    path = next((p for p in candidates if p.exists()), None)
    return ImageFont.truetype(str(path), size, index=1 if bold and path and path.suffix == ".ttc" else 0)


def make_flow(path: Path) -> None:
    img = Image.new("RGB", (1900, 470), "white")
    draw = ImageDraw.Draw(img)
    labels = ["医生问题与历史", "统一任务策略\nD1-D8 · 路由 · 证据", "专家与按需检索", "安全复核", "结构化终答"]
    colors = ["#FFF7DD", "#E8EEF5", "#EDF7F0", "#FDECEC", "#E8EEF5"]
    widths = [265, 380, 310, 240, 265]
    x, y, h = 60, 125, 210
    for i, (label, fill, w) in enumerate(zip(labels, colors, widths)):
        draw.rounded_rectangle((x, y, x + w, y + h), radius=24, fill=fill, outline="#8A94A6", width=3)
        lines = label.split("\n")
        for j, line in enumerate(lines):
            box = draw.textbbox((0, 0), line, font=font(38, bold=True))
            draw.text((x + (w - (box[2] - box[0])) / 2, y + 67 + j * 52), line, fill="#243447", font=font(38, bold=True))
        if i < len(labels) - 1:
            start = x + w + 14
            end = start + 50
            draw.line((start, y + h / 2, end, y + h / 2), fill="#667085", width=6)
            draw.polygon([(end, y + h / 2), (end - 18, y + h / 2 - 13), (end - 18, y + h / 2 + 13)], fill="#667085")
            x = end + 18
    img.save(path)


def make_cost_chart(path: Path) -> None:
    img = Image.new("RGB", (1800, 1000), "white")
    draw = ImageDraw.Draw(img)
    draw.text((80, 55), "专业月付：每 100 次会诊的技术成本结构（50 名付费用户）", fill="#243447", font=font(48, bold=True))
    items = [
        ("V4 Flash", 2.48, "#2E74B5"),
        ("Web Search", 1.92, "#7A9E7E"),
        ("支付费", 0.41, "#C69C48"),
        ("固定设施分摊", 24.00, "#8A94A6"),
        ("套餐收入", 69.00, "#C8754B"),
    ]
    x0, max_w, y0 = 420, 1180, 190
    for i, (name, value, color) in enumerate(items):
        y = y0 + i * 145
        draw.text((80, y + 25), name, fill="#243447", font=font(40, bold=True))
        w = max(8, int(max_w * value / 69.0))
        draw.rounded_rectangle((x0, y, x0 + w, y + 82), radius=16, fill=color)
        draw.text((min(x0 + w + 18, 1530), y + 18), f"¥{value:.2f}", fill="#243447", font=font(38, bold=True))
    draw.text((80, 925), "注：固定设施按 ¥1,200/月、50 名付费用户分摊；不含研发工资、税费与医疗责任成本。", fill="#667085", font=font(30))
    img.save(path)


def configure_weekly(doc: Document) -> None:
    existing = {style.name for style in doc.styles}
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
        if name not in existing:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if "Table Grid" not in existing:
        doc.styles.add_style("Table Grid", WD_STYLE_TYPE.TABLE)
    normal = doc.styles["Normal"]
    set_style(normal, size=11, color=BLACK, after=6, line=1.2)
    for name, size, before, after in (("Heading 1", 18, 18, 8), ("Heading 2", 16, 14, 6), ("Heading 3", 12, 10, 4)):
        set_style(doc.styles[name], size=size, color=BLACK, bold=True, before=before, after=after, line=1.2)
        doc.styles[name].paragraph_format.keep_with_next = True
    add_numbering(doc, abstract_id=177, num_id=177, fmt="bullet", text="•", left=720, hanging=360)


def build_weekly() -> Path:
    shutil.copy2(WEEKLY_REFERENCE, WEEKLY_OUTPUT)
    doc = Document(WEEKLY_OUTPUT)
    clear_body(doc)
    configure_weekly(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(6)
    title.paragraph_format.space_after = Pt(24)
    set_run_font(title.add_run("PetMind MoE Agent · 工作周报·8.10-8.16"), size=26, bold=True, color=BLACK)

    doc.add_heading("一、本周工作范围", level=1)
    add_callout(doc, "本周主线", "完成 PetMind 生产前端、统一 Agent 状态机与提示词策略，并形成可执行的服务定价和运营 SOP。", fill=LIGHT_GRAY, width=8280)
    add_table(doc, ["工作方向", "本周交付"], [
        ("新开发", "PetMind 生产前端：登录、套餐、聊天、历史会话、API Key 与后台管理"),
        ("架构优化", "将 D1-D8、专家路由和检索决策合并为一次统一任务策略"),
        ("质量验证", "24 个真实 API 混合病例全量通过；/chat-moe 与 /chat 完成真实联调"),
        ("运营准备", "重写服务定价、成本测算、上线巡检和事故处置 SOP"),
    ], [1900, 6380], header_fill=LIGHT_BLUE, font_size=10.5)

    doc.add_heading("二、工作内容简报", level=1)
    doc.add_heading("新增 PetMind 生产前端", level=2)
    add_bullet(doc, "完成兽医账号登录、邀请注册、套餐与订单、聊天会诊、历史搜索、会话导出和后台管理页面。", num_id=177)
    add_bullet(doc, "前端采用暖色纸张与铅绘风格，补齐加载、复制反馈、侧边栏收起、删除和导出等交互细节。", num_id=177)
    add_bullet(doc, "聊天采用 SSE 流式输出，刷新后可按 Run 事件续传，并从数据库恢复最终消息。", num_id=177)

    doc.add_heading("Agent 状态机与提示词优化", level=2)
    add_bullet(doc, "将原有多段意图识别、关键词判断和 Router 决策合并为一次统一任务策略，减少重复判断。", num_id=177)
    add_bullet(doc, "保留 D1-D8 的分类边界、路由指导和输出范式，让不同兽医任务仍按各自结构回答。", num_id=177)
    add_bullet(doc, "检索改为按证据需要触发：明确核验、最新资料和高风险用药必须检索，普通问题不强制调用工具。", num_id=177)

    doc.add_heading("SOP 报告与运营准备", level=2)
    add_bullet(doc, "基于真实 Token、工具调用和本地数据库快照重新测算 V4 Flash 与 Web Search 成本。", num_id=177)
    add_bullet(doc, "形成试用、月付、年付套餐建议，并补齐上线门禁、每日巡检、对账、降级和事故处置流程。", num_id=177)

    doc.add_heading("三、具体工作、开发与改动", level=1)
    doc.add_heading("改动 1：生产级网页端落地", level=2)
    add_body(doc, "现象：原有测试页面面向内部调试，缺少正式账号、套餐、会话资产和运营后台，无法直接提供对外服务。", bold_prefix="现象：")
    add_body(doc, "处理：新增独立 React 前端与专用后端路由，接入 JWT、API Key、角色权限、PostgreSQL 会话、积分和订单模型；对话运行先持久化再执行。", bold_prefix="处理：")
    add_body(doc, "结果：形成从登录、付费、聊天到管理审计的完整产品入口；会话、消息、运行事件和用户权限均可追溯。", bold_prefix="结果：")

    doc.add_heading("改动 2：统一 Agent 决策链路", level=2)
    flow_path = WORK / "weekly-unified-flow.png"
    make_flow(flow_path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(flow_path), width=Inches(5.0))
    set_alt_text(shape, "统一任务策略流程", "医生问题经过统一任务策略、专家与按需检索、安全复核后形成结构化终答。")
    cap = doc.add_paragraph("图 1  当前 MoE 主链路")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(cap.runs[0], size=9.5, color=MUTED)
    add_body(doc, "现象：旧链路中意图分类、关键词检索和 Router 分别决策，存在重复调用、边界不一致和维护成本高的问题。", bold_prefix="现象：")
    add_body(doc, "处理：统一策略一次输出主意图、次意图、专家、急症判断和证据任务；程序只做确定性校验，专家按 required/recommended 执行工具。", bold_prefix="处理：")
    add_body(doc, "结果：24 个真实 DeepSeek 混合病例全部通过；RAG 命中 10 题、Web Search 成功 8 题、7 题正常无工具，未出现为检索而检索。", bold_prefix="结果：")

    doc.add_heading("改动 3：服务成本与运营闭环", level=2)
    add_body(doc, "现象：旧报告基于过期峰谷价和较小样本，无法反映 V4 Flash 当前单价与统一策略的真实 Token。", bold_prefix="现象：")
    add_body(doc, "处理：按官方最新人民币价格、24 题共 449,000 Token、Tavily 当前信用点价格和固定设施假设重新测算，并区分模型、搜索与固定成本。", bold_prefix="处理：")
    add_body(doc, "结果：建议首期继续邀请制，采用试用 30 积分、月付 ¥69/600、年付 ¥699/7,200；以真实 usage、缓存命中和工具账单做月度复核。", bold_prefix="结果：")

    doc.add_heading("四、验证结果与下一步", level=1)
    add_table(doc, ["验证项", "结果"], [
        ("MoE + MCP 自动回归", "203 passed"),
        ("D1-D8 真实混合病例", "24/24；总计 449,000 Token"),
        ("工具执行", "RAG 10 题、Web 8 题、正常无工具 7 题"),
        ("/chat-moe", "两轮会话恢复成功；记住首轮宠物名；专家上下文恢复"),
        ("/chat", "页面、API Key、SSE 和 stop 事件正常"),
    ], [2600, 5680], header_fill=LIGHT_BLUE, font_size=10.5)
    add_body(doc, "下一步：在服务器采集 7 天真实 usage 与队列数据，完成支付对账、告警阈值和长期记忆画像的生产验收。", bold_prefix="下一步：")

    doc.core_properties.title = "PetMind MoE Agent 工作周报 8.10-8.16"
    doc.core_properties.subject = "生产前端、统一状态机、提示词与运营 SOP"
    doc.core_properties.author = "PetMind"
    doc.save(WEEKLY_OUTPUT)
    return WEEKLY_OUTPUT


def configure_sop(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    set_style(doc.styles["Normal"], size=11, color=INK, after=6, line=1.1)
    set_style(doc.styles["Title"], size=28, color=DARK_BLUE, bold=True, before=0, after=10, line=1.0)
    set_style(doc.styles["Subtitle"], size=13, color=MUTED, before=0, after=18, line=1.1)
    set_style(doc.styles["Heading 1"], size=16, color=BLUE, bold=True, before=16, after=8, line=1.1)
    set_style(doc.styles["Heading 2"], size=13, color=BLUE, bold=True, before=12, after=6, line=1.1)
    set_style(doc.styles["Heading 3"], size=12, color=DARK_BLUE, bold=True, before=8, after=4, line=1.1)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].paragraph_format.keep_with_next = True
    add_numbering(doc, abstract_id=77, num_id=77, fmt="bullet", text="•", left=720, hanging=360)
    add_numbering(doc, abstract_id=78, num_id=78, fmt="decimal", text="%1.", left=720, hanging=360)

    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(hp.add_run("PetMind · 服务定价与运营 SOP"), size=9, color=MUTED)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(fp.add_run("PetMind  |  2026-08-16  |  "), size=9, color=MUTED)
    add_field(fp, "PAGE")


def add_cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(kicker.add_run("PETMIND · RESEARCH & OPERATIONS"), size=11, bold=True, color=BLUE)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("MoE 服务定价与运营 SOP\n调研报告")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("统一任务策略 · 真实 API 回归 · DeepSeek V4 Flash 成本 · 套餐建议")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(meta.add_run("调研日期  2026-08-16    |    版本  v1.2"), size=10.5, color=MUTED)
    for _ in range(4):
        doc.add_paragraph()
    add_callout(doc, "一句话建议", "先以邀请制封闭试运营上线：试用 30 积分/7 天，专业月付 ¥69/600 积分，专业年付 ¥699/7,200 积分；先采集 7 天真实账单再扩大流量。", fill=PALE_GOLD)
    doc.add_page_break()


def build_sop() -> Path:
    doc = Document()
    configure_sop(doc)
    add_cover(doc)

    doc.add_heading("管理层摘要", level=1)
    add_table(doc, ["观察项", "真实结果", "经营含义"], [
        ("统一策略回归", "D1-D8 各 3 题，24/24 通过", "新链路可替代旧分类器、关键词策略和 Router LLM"),
        ("真实 Token", "24 题共 449,000；均值 18,708/题", "复杂 MoE 是多次内部 LLM 调用，不能按单次问答估价"),
        ("检索执行", "RAG 10 题；Web 8 题；正常无工具 7 题", "工具率应按证据需要解释，不追求全量调用"),
        ("双入口", "/chat-moe 两轮恢复；/chat SSE 正常", "测试会话与无状态兼容接口均可用"),
        ("变量成本", "典型约 ¥0.044/题（含当前 Web 结构）", "模型便宜，固定设施、运维和专业责任是定价主体"),
        ("本机 QA 快照", "100 条；98 条 MoE；平均 30.91 s", "历史样本仍需逐步补齐真实 usage 与工具审计"),
    ], [1600, 2800, 4960], header_fill=LIGHT_BLUE, font_size=9.2)
    add_callout(doc, "Go / No-Go", "当前适合小规模邀请制试运营，不适合公开无限量售卖。Go 条件：真实 usage 入库、支付与退款幂等、急症基准、跨用户隔离和 7 天账单对齐全部通过。", fill=PALE_RED)

    doc.add_heading("1. 调研范围与事实口径", level=1)
    add_body(doc, "本报告从头重算 PetMind MoE 的模型、搜索和固定设施成本，并把统一任务策略、数据库持久化、双入口联调和运营动作放进同一份 SOP。")
    doc.add_heading("1.1 数据来源", level=2)
    add_bullet(doc, "24 个真实 DeepSeek API 混合病例：D1-D8 各 3 题，保留分类、工具、终答和总 Token。")
    add_bullet(doc, "本机 QA SQLite：100 条历史记录，其中 98 条 agent-moe；4 条 Web Search，3 条 RAG 命中。")
    add_bullet(doc, "本机 Session SQLite：4 个测试会话；成功联调会话保存 2 轮消息与 2 轮专家上下文。")
    add_bullet(doc, "DeepSeek 官方价格、缓存、并发文档；Tavily 官方信用点与 Search 计费文档。")
    doc.add_heading("1.2 计算边界", level=2)
    add_bullet(doc, "V4 Flash 使用官方人民币价：缓存命中 ¥0.02/M、未命中 ¥1/M、输出 ¥2/M Token。")
    add_bullet(doc, "24 题结果只保存 total_tokens，输入/输出拆分用旧基准的 67.5%/32.5%估算；同时给出保守上界。")
    add_bullet(doc, "Tavily Basic Search 按 1 credit/次、$0.008/credit 和内部规划汇率 ¥7.20/USD 测算。")
    add_bullet(doc, "固定设施按 ¥1,200/月假设；不包含研发工资、税费、发票、退款和医疗责任成本。")

    doc.add_heading("2. 当前系统与验证结果", level=1)
    doc.add_heading("2.1 生产链路", level=2)
    add_table(doc, ["阶段", "当前职责", "关键控制"], [
        ("统一任务策略", "一次输出 D1-D8、专家、急症与证据任务", "完整分类边界和 routing_guidance；无关键词判断"),
        ("确定性门控", "校验范围、工具可用性与 owner", "不增加第二次 Router LLM"),
        ("专家执行", "临床/药学/营养/行为并行分析", "required 必须完成；recommended 不阻止 final"),
        ("Critic", "事实与安全复核", "可限制危险内容，不得绕过 D1-D8 输出契约"),
        ("Aggregator", "按主意图组织终答", "只引用本轮真实工具结果"),
    ], [1500, 3300, 4560], header_fill=LIGHT_BLUE, font_size=9.2)
    doc.add_heading("2.2 24 题真实回归", level=2)
    add_table(doc, ["维度", "题数", "Token", "RAG", "Web", "无工具"], [
        ("D1", 3, "45,146", 1, 1, 1), ("D2", 3, "50,811", 1, 1, 1),
        ("D3", 3, "48,684", 1, 1, 1), ("D4", 3, "53,117", 1, 1, 1),
        ("D5", 3, "65,065", 2, 1, 1), ("D6", 3, "59,837", 2, 1, 0),
        ("D7", 3, "66,944", 1, 1, 1), ("D8", 3, "59,396", 1, 1, 1),
        ("合计", 24, "449,000", 10, 8, 7),
    ], [1200, 900, 1900, 1300, 1300, 1760], header_fill=LIGHT_GRAY, font_size=9.2)
    add_body(doc, "注：检索类别可重叠；有 1 个病例在本地证据不足时继续使用 Web 补证，因此各列不能简单相加为 24。")
    add_callout(doc, "验收结果", "意图、输出变体、D1-D8 结构、required 工具闭环与无工具边界均为 24/24；MoE + MCP 自动回归 203 passed。", fill=PALE_GREEN)

    doc.add_heading("2.3 /chat-moe 与 /chat 对接", level=2)
    add_table(doc, ["入口", "验证动作", "结果"], [
        ("/chat-moe", "创建 Session，连续发送两轮问题", "第二轮加载 1 个完整轮次与 1 个专家上下文，正确回忆猫名"),
        ("/chat", "页面调用 /v1/chat/completions", "API Key 鉴权、SSE、终止事件正常，无 error/busy"),
        ("会话原子性", "上游失败后检查 SQLite", "失败轮次不提交；成功轮次保存消息与专家上下文"),
    ], [1500, 3100, 4760], header_fill=LIGHT_BLUE, font_size=9.4)
    add_body(doc, "部署注意：后端必须能访问上游模型。系统代理不可用时设置 HTTPX_TRUST_ENV=0；这影响出站 LLM 连接，不影响页面路由本身。")

    doc.add_heading("3. DeepSeek V4 Flash 与搜索成本", level=1)
    doc.add_heading("3.1 官方当前价格", level=2)
    add_table(doc, ["项目", "V4 Flash 单价", "说明"], [
        ("输入 · 缓存命中", "¥0.02 / 百万 Token", "DeepSeek 磁盘缓存自动启用，按实际命中量计费"),
        ("输入 · 缓存未命中", "¥1 / 百万 Token", "套餐测算采用该保守口径"),
        ("输出", "¥2 / 百万 Token", "通常是单次会诊模型成本的主要部分"),
        ("账户并发", "2,500", "首期瓶颈仍在本地单 GPU Worker 和 MoE 串行阶段"),
    ], [2300, 2200, 4860], header_fill=LIGHT_GRAY, font_size=9.5)
    add_body(doc, "官方同时说明 deepseek-chat/deepseek-reasoner 为兼容别名；工程默认和部署示例已显式切换为 deepseek-v4-flash，避免依赖旧名。")
    doc.add_heading("3.2 24 题成本推算", level=2)
    add_table(doc, ["口径", "每题 Token/调用", "模型成本", "含 Web 后"], [
        ("典型估算", "输入 12,628；输出 6,080", "约 ¥0.0248", "按 8/24 Web 占比，约 ¥0.0440/题"),
        ("50% 输入缓存命中", "输入各半；输出不变", "约 ¥0.0186", "约 ¥0.0378/题"),
        ("保守上界", "18,708 全按输出价", "约 ¥0.0374", "单题如执行 1 次 Basic Search，再加约 ¥0.0576"),
    ], [2100, 3000, 1900, 2360], header_fill=LIGHT_BLUE, font_size=9.4)
    add_body(doc, "24 题整批模型成本典型估算约 ¥0.60；8 次 Basic Search 约 ¥0.46；合计约 ¥1.06。RAG 为本地服务，不产生按次供应商费用，但会占用 CPU/GPU、内存和运维预算。")

    cost_path = WORK / "sop-cost-structure.png"
    make_cost_chart(cost_path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(cost_path), width=Inches(6.35))
    set_alt_text(shape, "专业月付成本结构图", "50 名付费用户时，100 次会诊的模型、Web、支付费、固定设施分摊与套餐收入对比。")
    cap = doc.add_paragraph("图 1  专业月付的技术成本结构")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(cap.runs[0], size=9.5, color=MUTED)

    doc.add_heading("4. 推荐定价与套餐", level=1)
    add_table(doc, ["套餐", "价格/有效期", "积分", "约合会诊", "建议权益"], [
        ("试用 trial", "¥0 / 7 天", "30", "约 5 次", "聊天、历史、导出；不开放 API Key"),
        ("专业月付", "¥69 / 30 天", "600", "约 100 次", "聊天、搜索、记忆、1 个 API Key、5 req/min"),
        ("专业年付", "¥699 / 365 天", "7,200", "约 1,200 次", "月付全部权益、2 个 API Key；年化优惠约 15.6%"),
        ("加量包", "¥19", "150", "约 25 次", "仅活跃会员可购，随订阅到期"),
    ], [1300, 1650, 1000, 1450, 3960], header_fill=LIGHT_BLUE, font_size=9.2)
    add_body(doc, "积分采用 6 积分/标准 MoE 会诊的简单口径。积分表示容量与复杂度配额，不直接等同 Token 或人民币；失败且没有有效模型用量时退回预占。")
    doc.add_page_break()
    doc.add_heading("4.1 月付盈亏", level=2)
    add_table(doc, ["付费用户", "固定成本/人", "变量成本/人", "支付费/人", "合计/人", "技术毛贡献率"], [
        (25, "¥48.00", "¥4.40", "¥0.41", "¥52.81", "23.5%"),
        (50, "¥24.00", "¥4.40", "¥0.41", "¥28.81", "58.2%"),
        (100, "¥12.00", "¥4.40", "¥0.41", "¥16.81", "75.6%"),
    ], [1200, 1700, 1700, 1500, 1500, 1760], header_fill=LIGHT_GRAY, font_size=9.2)
    add_callout(doc, "经营阈值", "在上述假设下，月付固定成本盈亏平衡约 19 名付费用户。低于 25 名时优先采用封闭试点或机构赞助，不推出低价无限量。", fill=PALE_GOLD)
    add_body(doc, "定价不是 Token 成本乘倍数。用户购买的是兽医知识编排、可追溯证据、持续更新、跨会话资产、稳定队列与人工支持；医疗责任和服务保障必须留出预算。")

    doc.add_heading("5. 上线与运营 SOP", level=1)
    doc.add_heading("5.1 上线前门禁", level=2)
    for text in [
        "模型显式设置 deepseek-v4-flash；启动日志只记录 model/base_url，不记录 API Key。",
        "每个内部 LLM 调用保存输入、输出、缓存命中/未命中 Token；Run 汇总供应商成本与积分结算。",
        "完成空库/已有库迁移、API Key 撤销、JWT 轮换、跨用户隔离、订单回调幂等和积分并发预占测试。",
        "执行 203 项 MoE/MCP 回归与 24 题真实门禁；任何急症分级错误、跨用户泄漏或结构契约失败均阻止发布。",
        "验证 /chat、/chat-moe、生产 conversations/runs、SSE 断线续传、Redis/Worker 重启恢复和 Memory 降级。",
    ]:
        add_number(doc, text)

    doc.add_heading("5.2 每日巡检", level=2)
    add_table(doc, ["检查项", "告警阈值", "动作"], [
        ("API、Worker、Memory、PostgreSQL、Redis", "任一不健康或 /ready=false", "停止新任务入队；保留已持久化 Run；恢复后重放"),
        ("Run 成功率与排队", "成功率 <99%；P95 >90s；队列 >20", "检查上游、专家数与工具；限流或降级"),
        ("DeepSeek", "任何 401；连续 3 次 429；余额 <7 天", "401 不盲重试；切换 Key；429 指数退避"),
        ("Token/成本", "单次 >30k；日成本超预算 20%", "定位历史膨胀、提示词回退或工具循环"),
        ("积分一致性", "预占 >30min；余额不平", "幂等补偿；失败无有效用量则退款"),
        ("Memory", "消息与回执不等；失败率 >1%", "主链继续，异步重试；禁止重复 turn_id"),
    ], [2600, 2800, 3960], header_fill=LIGHT_BLUE, font_size=9.2)

    doc.add_heading("5.3 每周与每月", level=2)
    add_bullet(doc, "每周抽检 20 条：覆盖宠物主/兽医、急症、用药、营养、D1-D8、RAG、Web 和无工具题。")
    add_bullet(doc, "每周统计 task_policy 准确率、专家数、required/recommended、工具成功率、缓存命中率、Token P50/P95 和重试率。")
    add_bullet(doc, "每月复核 DeepSeek/Tavily 官方价格和模型公告，用真实账单回算套餐毛利；价格变化 >10% 时重算。")
    add_bullet(doc, "每月执行 PostgreSQL 备份恢复、Redis 丢失恢复、过期 Token/邀请清理和跨用户隔离抽测。")

    doc.add_heading("5.4 事故处置", level=2)
    add_table(doc, ["事件", "立即处置", "结算原则"], [
        ("DeepSeek 401", "停用失效 Key、切换备用 Key、检查模型名", "无有效输出则全额退积分"),
        ("429 / 排队过长", "指数退避，保持 Run 持久化并发送排队阶段", "未开始模型调用不扣费"),
        ("5xx / 网络中断", "有限重试；SSE 按 Last-Event-ID 恢复", "用 run_id 去重，避免双结算"),
        ("RAG / Web / Memory 不可用", "降级并明确无法核实，不伪造来源", "工具失败不额外收费"),
        ("疑似数据泄漏", "撤销 Token/API Key，冻结审计，停止受影响接口", "进入安全事件与合规流程"),
    ], [1900, 4560, 2900], header_fill=LIGHT_GRAY, font_size=9.2)

    doc.add_page_break()
    doc.add_heading("6. 计量、工具与质量规则", level=1)
    add_table(doc, ["场景", "证据要求", "是否必须调用"], [
        ("用户明确要求检索、联网、最新版本或引用", "创建 required evidence task", "必须；失败要说明无法核实"),
        ("具体剂量、相互作用、禁忌、物种毒性、停换药", "药学资料；必要时 Web 补证", "高风险具体结论必须"),
        ("被动心率告警或真实患者体征", "vitals/SQL/MCP 个体事实", "必须；外部 flag 不是诊断事实"),
        ("常规诊断、报告解读、低风险知识问答", "recommended 或无证据任务", "不锁定；专家可直接 final"),
        ("tools=[] / tool_choice=none", "记录 unavailable 与局限", "调用方禁用优先，不绕过"),
    ], [3000, 3500, 2860], header_fill=LIGHT_BLUE, font_size=9.2)
    add_body(doc, "判断工具率时必须分层。当前 24 题中 7 题正常无工具是预期行为；工具调用成功只证明链路完成，不等同于医学结论正确，仍需做数字—来源和引用—工具结果审计。")

    doc.add_heading("7. 接口与部署要点", level=1)
    add_table(doc, ["接口", "状态管理", "适用场景"], [
        ("conversations/{id}/runs", "平台持久化会话、消息、Run、事件、积分与 Memory", "正式网页和需要断线恢复的客户"),
        ("/v1/chat/completions", "无状态；调用方重发 messages", "OpenAI 兼容的第三方系统"),
        ("/chat-moe/completions", "SQLite 测试会话与专家上下文", "内部验收，不作为新平台生产数据源"),
    ], [3000, 3700, 2660], header_fill=LIGHT_GRAY, font_size=9.2)
    add_bullet(doc, "系统级服务适合开机启动、统一日志和服务器运维；用户级服务无需 root，但依赖用户会话与 linger 配置。不要同时启用两套相同端口。")
    add_bullet(doc, "生产接入反向代理时设置可信 Host/CORS、上传和请求体上限、SSE 超时与 keep-alive；公网穿透的 TCP/RTT 延迟不能误判为 Agent 推理慢。")
    add_bullet(doc, "默认热启动 RAG、BM25、Reranker、分类索引和 Memory Embedding；只在 /ready=true 后接流量。")

    doc.add_page_break()
    doc.add_heading("8. 决策清单", level=1)
    add_table(doc, ["优先级", "事项", "完成标准"], [
        ("P0", "真实 usage 与供应商成本入库", "每个 Run 的 input/output/cache/provider_cost 非 0；可按日对账"),
        ("P0", "正式套餐与支付幂等", "¥69/600、¥699/7,200 可售；重复回调不重复发积分"),
        ("P0", "安全与质量门禁", "24 题、急症、跨用户隔离、API Key 撤销均通过"),
        ("P1", "容量与告警", "队列、P95、401/429/5xx、余额、Token 异常有告警"),
        ("P1", "长期记忆验收", "profile/page 生成、检索召回和数据隔离有量化测试"),
        ("P2", "提示词与 Thinking 成本优化", "质量不降前提下 Token/耗时下降至少 15%"),
    ], [1200, 3000, 5160], header_fill=LIGHT_BLUE, font_size=9.2)

    doc.add_heading("附录 A：公式与来源", level=1)
    add_body(doc, "成本公式：单次成本 = 输入缓存命中 Token × ¥0.02/M + 输入未命中 Token × ¥1/M + 输出 Token × ¥2/M + Search credits × 单价。")
    add_bullet(doc, "DeepSeek 模型与价格：https://api-docs.deepseek.com/zh-cn/quick_start/pricing（访问：2026-08-16）")
    add_bullet(doc, "DeepSeek 缓存：https://api-docs.deepseek.com/guides/kv_cache（访问：2026-08-16）")
    add_bullet(doc, "DeepSeek 限速与隔离：https://api-docs.deepseek.com/quick_start/rate_limit（访问：2026-08-16）")
    add_bullet(doc, "Tavily Credits & Pricing：https://docs.tavily.com/documentation/api-credits（访问：2026-08-16）")
    add_bullet(doc, "本地真实回归：agent_api/tests/moe/reports/unified_policy_regression_full_final_20260816/（24 题原始答复与断言，本地审计目录）")
    add_callout(doc, "复核要求", "供应商价格会变化。每月必须重新打开官方页面并用真实账单复算；本报告数字不是长期固定报价。", fill=PALE_RED)

    doc.core_properties.title = "PetMind MoE 服务定价与运营 SOP 调研报告 v1.2"
    doc.core_properties.subject = "真实 Token、统一策略、DeepSeek V4 Flash 成本与套餐建议"
    doc.core_properties.author = "PetMind"
    doc.save(SOP_OUTPUT)
    return SOP_OUTPUT


def main() -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    weekly = build_weekly()
    sop = build_sop()
    print(weekly)
    print(sop)


if __name__ == "__main__":
    main()
